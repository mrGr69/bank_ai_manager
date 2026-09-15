from __future__ import annotations

import os
import shutil
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.config import settings

KYIV = ZoneInfo(settings.tz)


def parse_money(value) -> Decimal | None:
    if value is None or value == "" or value == "—":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).replace("\xa0", " ").replace(" ", "").replace(",", ".")
    text = text.replace("UAH", "").strip()
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def parse_dt(value) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        dt = None
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=KYIV)
    return dt


def iter_xlsx_rows(path: Path):
    tmp: Path | None = None
    load_path = path
    if path.suffix.lower() != ".xlsx":
        fd, name = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        shutil.copy2(path, name)
        tmp = Path(name)
        load_path = tmp
    try:
        wb = load_workbook(load_path, data_only=True, read_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            yield row
        wb.close()
    finally:
        if tmp is not None:
            tmp.unlink(missing_ok=True)


def card_last4(mask: str | None) -> str:
    if not mask:
        return ""
    digits = "".join(ch for ch in str(mask) if ch.isdigit())
    return digits[-4:] if len(digits) >= 4 else ""
